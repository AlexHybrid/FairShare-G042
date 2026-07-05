import os
import sys

# Ensure python-pptx is installed
try:
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN
except ImportError:
    import subprocess
    subprocess.check_call([sys.executable, "-m", "pip", "install", "python-pptx"])
    from pptx import Presentation
    from pptx.util import Inches, Pt
    from pptx.dml.color import RGBColor
    from pptx.enum.text import PP_ALIGN

def create_presentation():
    prs = Presentation()

    # Define colors
    navy_bg = RGBColor(0x11, 0x18, 0x33)      # #111833
    white_text = RGBColor(0xFF, 0xFF, 0xFF)   # #FFFFFF
    slate_gray = RGBColor(0x94, 0xA3, 0xB8)   # #94a3b8
    accent_green = RGBColor(0x4A, 0xDE, 0x80) # #4ade80
    accent_rose = RGBColor(0xFB, 0x71, 0x85)  # #fb7185
    accent_blue = RGBColor(0x4F, 0x6B, 0xFF)  # #4f6bff

    # Use a blank layout (index 6 is blank usually)
    blank_slide_layout = prs.slide_layouts[6]

    def add_custom_slide(title_text, subtitle_text=None):
        slide = prs.slides.add_slide(blank_slide_layout)

        # Apply dark navy background
        background = slide.background
        fill = background.fill
        fill.solid()
        fill.fore_color.rgb = navy_bg

        # Add Slide Title
        txBox = slide.shapes.add_textbox(Inches(0.5), Inches(0.5), Inches(9.0), Inches(0.8))
        tf = txBox.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Segoe UI"
        p.font.size = Pt(36)
        p.font.bold = True
        p.font.color.rgb = white_text

        # Add Subtitle if provided
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Segoe UI"
            p2.font.size = Pt(14)
            p2.font.color.rgb = slate_gray

        return slide

    # --- Slide 1: Title Slide ---
    slide = prs.slides.add_slide(blank_slide_layout)
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = navy_bg

    # Title & Subtitle in one text frame
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.2), Inches(9.0), Inches(3.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Rental Blossoms"
    p.font.name = "Segoe UI"
    p.font.size = Pt(54)
    p.font.bold = True
    p.font.color.rgb = accent_green
    p.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "FairShare: Smart Co-Living & Bill Splitter"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(22)
    p2.font.color.rgb = white_text
    p2.alignment = PP_ALIGN.CENTER

    p3 = tf.add_paragraph()
    p3.text = "\nFull Technical System Walkthrough & Architecture Presentation"
    p3.font.name = "Segoe UI"
    p3.font.size = Pt(14)
    p3.font.italic = True
    p3.font.color.rgb = slate_gray
    p3.alignment = PP_ALIGN.CENTER

    # --- Slide 2: The Core Problem & Solution ---
    slide = add_custom_slide("The Problem & The Solution", "Why Rental Blossoms (FairShare) exists")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    # Problem section
    p = tf.paragraphs[0]
    p.text = "THE PROBLEM:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_rose

    bullet1 = tf.add_paragraph()
    bullet1.text = "• Manual Split calculations for shared bills (electricity, water, wifi) are error-prone."
    bullet1.font.size = Pt(16)
    bullet1.font.color.rgb = white_text
    bullet1.level = 1

    bullet2 = tf.add_paragraph()
    bullet2.text = "• Receipts get lost in messy chat histories, leaving no clear audit trail."
    bullet2.font.size = Pt(16)
    bullet2.font.color.rgb = white_text
    bullet2.level = 1

    # Solution section
    p2 = tf.add_paragraph()
    p2.text = "\nTHE SOLUTION: RENTAL BLOSSOMS (FAIRSHARE)"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(18)
    p2.font.bold = True
    p2.font.color.rgb = accent_green

    bullet3 = tf.add_paragraph()
    bullet3.text = "• Real-time digital dashboard representing all expenses dynamically."
    bullet3.font.size = Pt(16)
    bullet3.font.color.rgb = white_text
    bullet3.level = 1

    bullet4 = tf.add_paragraph()
    bullet4.text = "• Secure receipt upload and persistence linked to database records."
    bullet4.font.size = Pt(16)
    bullet4.font.color.rgb = white_text
    bullet4.level = 1

    bullet5 = tf.add_paragraph()
    bullet5.text = "• Automatic PDF and Excel sheet generation for quick auditing."
    bullet5.font.size = Pt(16)
    bullet5.font.color.rgb = white_text
    bullet5.level = 1

    # --- Slide 3: System Architecture (The Restaurant Analogy) ---
    slide = add_custom_slide("System Architecture Overview", "How the frontend, backend, and database interact")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Our software follows a clean Client-Server Model:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_blue

    b1 = tf.add_paragraph()
    b1.text = "• Frontend (The Restaurant Table): Built with clean HTML, Vanilla CSS, and JavaScript. Users see details and order actions."
    b1.font.size = Pt(15)
    b1.font.color.rgb = white_text
    b1.level = 1

    b2 = tf.add_paragraph()
    b2.text = "• HTTP Requests (The Waiter): Asynchronous JavaScript Fetch API queries act as the bridge, conveying JSON data over the network."
    b2.font.size = Pt(15)
    b2.font.color.rgb = white_text
    b2.level = 1

    b3 = tf.add_paragraph()
    b3.text = "• Flask Server (The Kitchen / app.py): Directs traffic, processes computations, verifies login sessions, and serves files."
    b3.font.size = Pt(15)
    b3.font.color.rgb = white_text
    b3.level = 1

    b4 = tf.add_paragraph()
    b4.text = "• SQLite Database (The Pantry / database.py): Keeps a persistent and structured ledger of users and expenses."
    b4.font.size = Pt(15)
    b4.font.color.rgb = white_text
    b4.level = 1

    # --- Slide 4: Database Layer & Smart Migrations ---
    slide = add_custom_slide("Database Layer: sqlite3 & database.py", "Robust local persistence with safety migrations")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Key SQLite & database.py Features:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_green

    b1 = tf.add_paragraph()
    b1.text = "• Row Factory Access: Uses conn.row_factory = sqlite3.Row. This allows developers to reference query results by column names (e.g. row['amount']) rather than hardcoded indexes (e.g. row[3])."
    b1.font.size = Pt(15)
    b1.font.color.rgb = white_text
    b1.level = 1

    b2 = tf.add_paragraph()
    b2.text = "• Auto-Migration System: On every startup, it runs PRAGMA table_info(expenses) to check structure. If update columns are missing, it appends them dynamically (e.g. ALTER TABLE) without wiping out user database records."
    b2.font.size = Pt(15)
    b2.font.color.rgb = white_text
    b2.level = 1

    b3 = tf.add_paragraph()
    b3.text = "• DB Seeding: Automatically pre-populates default mock data (Rooms A-E) if the db is fresh, allowing instant demo capability."
    b3.font.size = Pt(15)
    b3.font.color.rgb = white_text
    b3.level = 1

    # --- Slide 5: Flask Backend API ---
    slide = add_custom_slide("Flask Backend: app.py API", "Serving files, processing endpoints, and security")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Central Server Features (Port 5000):"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_rose

    b1 = tf.add_paragraph()
    b1.text = "• Session-based Authentication: Secure login/signup handlers using Flask's native session management and SHA-256 hashed password checks."
    b1.font.size = Pt(15)
    b1.font.color.rgb = white_text
    b1.level = 1

    b2 = tf.add_paragraph()
    b2.text = "• Asynchronous REST API Endpoints: Handles GET /api/expenses to load ledger entries and POST /api/expenses to save new ones."
    b2.font.size = Pt(15)
    b2.font.color.rgb = white_text
    b2.level = 1

    b3 = tf.add_paragraph()
    b3.text = "• Secure File Uploads: Employs secure_filename() to sanitize user-provided file names. Restricts uploads strictly to images (PNG, JPG) and PDF, preventing path-traversal scripts from hacking the host OS."
    b3.font.size = Pt(15)
    b3.font.color.rgb = white_text
    b3.level = 1

    # --- Slide 6: Real-Time UI & Fetch Requests ---
    slide = add_custom_slide("Frontend Integration: JS & CSS", "Making the dashboard alive, responsive, and dynamic")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "How JS & CSS Deliver a Premium User Experience:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_blue

    b1 = tf.add_paragraph()
    b1.text = "• Two-Step Network Handshake: Adding a bill does two things. First, it POSTs the physical receipt file to the server and retrieves the saved path. Second, it POSTs the JSON payload including that receipt path to the database. This separates heavy file traffic from structured records."
    b1.font.size = Pt(15)
    b1.font.color.rgb = white_text
    b1.level = 1

    b2 = tf.add_paragraph()
    b2.text = "• Chart.js Integration: Aggregates room-by-room totals client-side and renders a responsive, high-performance doughnut chart reflecting active split proportions."
    b2.font.size = Pt(15)
    b2.font.color.rgb = white_text
    b2.level = 1

    b3 = tf.add_paragraph()
    b3.text = "• Theme Styling: Custom theme selectors apply custom room tags and colors so tenants can identify their bills instantly."
    b3.font.size = Pt(15)
    b3.font.color.rgb = white_text
    b3.level = 1

    # --- Slide 7: Integrated Reporting Engines ---
    slide = add_custom_slide("Integrated Reporting Engines", "Compiling production reports in app.py")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Landlords can generate two rich file formats instantly:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_green

    b1 = tf.add_paragraph()
    b1.text = "• Excel Report Engine (openpyxl): Queries the SQLite database, builds a spreadsheet workbook, applies custom typography styles, formats the currency column as ('RM'#,##0.00), auto-resizes columns, and returns a downloadable .xlsx file."
    b1.font.size = Pt(15)
    b1.font.color.rgb = white_text
    b1.level = 1

    b2 = tf.add_paragraph()
    b2.text = "• PDF Report Engine (reportlab): Builds a structured document using flowable objects. Highlights header rows in navy blue, aggregates total transaction values programmatically, handles tabular page layout margins, and serves a crisp vector printout."
    b2.font.size = Pt(15)
    b2.font.color.rgb = white_text
    b2.level = 1

    # --- Slide 8: Standalone Microservices ---
    slide = add_custom_slide("Standalone Export Microservices", "Decoupled reporting modules for high scale")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(1.5), Inches(9.0), Inches(5.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Isolated Microservices:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(18)
    p.font.bold = True
    p.font.color.rgb = accent_rose

    b1 = tf.add_paragraph()
    b1.text = "• Excel Service (ExcelExport.py - Port 5001): Standalone Flask instance simulating specialized Excel compile operations."
    b1.font.size = Pt(15)
    b1.font.color.rgb = white_text
    b1.level = 1

    b2 = tf.add_paragraph()
    b2.text = "• PDF Service (pdfExport.py - Port 5002): Standalone Flask instance compiling PDF document formats on an isolated process."
    b2.font.size = Pt(15)
    b2.font.color.rgb = white_text
    b2.level = 1

    b3 = tf.add_paragraph()
    b3.text = "• Architectural Value: Shows how the application can scale. Offloading heavy file processing to separate ports prevents the core Flask API server (Port 5000) from slowing down during high-traffic landlord audits."
    b3.font.size = Pt(15)
    b3.font.color.rgb = white_text
    b3.level = 1

    # --- Slide 9: Summary & Q&A ---
    slide = add_custom_slide("Summary & Questions", "Thank you for listening!")
    txBox = slide.shapes.add_textbox(Inches(0.5), Inches(2.0), Inches(9.0), Inches(4.0))
    tf = txBox.text_frame
    tf.word_wrap = True

    p = tf.paragraphs[0]
    p.text = "Rental Blossoms (FairShare) Key Strengths:"
    p.font.name = "Segoe UI"
    p.font.size = Pt(20)
    p.font.bold = True
    p.font.color.rgb = accent_green
    p.alignment = PP_ALIGN.CENTER

    b1 = tf.add_paragraph()
    b1.text = "✓ Modern, responsive UI with real-time Chart.js graphics"
    b1.font.size = Pt(16)
    b1.font.color.rgb = white_text
    b1.alignment = PP_ALIGN.CENTER

    b2 = tf.add_paragraph()
    b2.text = "✓ Resilient database layer with built-in safety columns migrations"
    b2.font.size = Pt(16)
    b2.font.color.rgb = white_text
    b2.alignment = PP_ALIGN.CENTER

    b3 = tf.add_paragraph()
    b3.text = "✓ Multi-format reporting architecture with microservice compatibility"
    b3.font.size = Pt(16)
    b3.font.color.rgb = white_text
    b3.alignment = PP_ALIGN.CENTER

    p2 = tf.add_paragraph()
    p2.text = "\n\nAny Questions?"
    p2.font.name = "Segoe UI"
    p2.font.size = Pt(28)
    p2.font.bold = True
    p2.font.color.rgb = accent_blue
    p2.alignment = PP_ALIGN.CENTER

    # Save presentation
    output_path = os.path.join(os.path.dirname(__file__), "rental_blossoms_presentation.pptx")
    prs.save(output_path)
    print(f"Presentation saved successfully to: {output_path}")

if __name__ == "__main__":
    create_presentation()